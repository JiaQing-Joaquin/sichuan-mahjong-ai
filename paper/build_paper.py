"""Build the Chinese technical report (DOCX) for the Sichuan Mahjong AI."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
FIG = ROOT / "paper" / "figures"
OUT = ROOT / "paper" / "川麻血战到底AI技术报告.docx"

BODY_FONT = "宋体"
HEAD_FONT = "黑体"
LATIN_FONT = "Times New Roman"


def set_run_font(run, name: str, size: float | None = None, bold: bool | None = None,
                 color: RGBColor | None = None):
    run.font.name = LATIN_FONT
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = color


def body(text: str, indent: bool = True, size: float = 10.5, space_after: float = 4):
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_after = Pt(space_after)
    if indent:
        p.paragraph_format.first_line_indent = Cm(0.74)
    run = p.add_run(text)
    set_run_font(run, BODY_FONT, size)
    return p


def h1(text: str, number: str | None = None):
    p = doc.add_paragraph(style="Heading 1")
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run((f"{number}  " if number else "") + text)
    set_run_font(run, HEAD_FONT, 15, bold=True, color=RGBColor(0, 0, 0))
    return p


def h2(text: str, number: str | None = None):
    p = doc.add_paragraph(style="Heading 2")
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run((f"{number}  " if number else "") + text)
    set_run_font(run, HEAD_FONT, 12, bold=True, color=RGBColor(0, 0, 0))
    return p


def caption(text: str, above: bool = False):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6 if not above else 8)
    p.paragraph_format.space_after = Pt(6 if above else 8)
    p.paragraph_format.keep_with_next = above
    run = p.add_run(text)
    set_run_font(run, BODY_FONT, 9, color=RGBColor(64, 64, 64))


def figure(name: str, cap: str, width_cm: float = 15.0):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(str(FIG / name), width=Cm(width_cm))
    caption(cap)


def shade(cell, color: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), color)
    tc_pr.append(shd)


def table(headers, rows, widths=None, font_size: float = 9):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(headers):
        cell = t.rows[0].cells[i]
        cell.text = ""
        para = cell.paragraphs[0]
        para.paragraph_format.line_spacing = 1.0
        para.paragraph_format.space_before = Pt(0)
        para.paragraph_format.space_after = Pt(0)
        run = para.add_run(h)
        set_run_font(run, HEAD_FONT, font_size, bold=True)
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        shade(cell, "DEEAF6")
    for row in rows:
        cells = t.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = ""
            para = cells[i].paragraphs[0]
            para.paragraph_format.line_spacing = 1.0
            para.paragraph_format.space_before = Pt(0)
            para.paragraph_format.space_after = Pt(0)
            run = para.add_run(str(value))
            set_run_font(run, BODY_FONT, font_size)
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if widths:
        for row in t.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Cm(w)

    for row in t.rows:
        row.height = Cm(0.6)
        row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST

    # repeat the header row on page breaks and keep the table on one page
    tr_pr = t.rows[0]._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)
    for row in t.rows[:-1]:
        for cell in row.cells:
            for p in cell.paragraphs:
                p.paragraph_format.keep_with_next = True

    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_page_number(section):
    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    set_run_font(run, BODY_FONT, 9, color=RGBColor(96, 96, 96))
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_end)


doc = Document()
section = doc.sections[0]
section.page_width = Cm(21.0)
section.page_height = Cm(29.7)
section.top_margin = Cm(2.4)
section.bottom_margin = Cm(2.4)
section.left_margin = Cm(2.5)
section.right_margin = Cm(2.5)

style = doc.styles["Normal"]
style.font.name = LATIN_FONT
style.font.size = Pt(10.5)
style.element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)

add_page_number(section)

# ---------------------------------------------------------------------------
# Title
# ---------------------------------------------------------------------------
title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.paragraph_format.space_before = Pt(6)
title.paragraph_format.space_after = Pt(4)
run = title.add_run("四川麻将「血战到底」AI：规则引擎、监督学习与自博弈实践")
set_run_font(run, HEAD_FONT, 18, bold=True)

subtitle = doc.add_paragraph()
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
subtitle.paragraph_format.space_after = Pt(2)
run = subtitle.add_run("Sichuan Mahjong “Blood Battle to the End” AI: "
                       "Rule Engine, Imitation Learning and Self-Play")
set_run_font(run, BODY_FONT, 10.5, color=RGBColor(64, 64, 64))

meta = doc.add_paragraph()
meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
meta.paragraph_format.space_after = Pt(10)
run = meta.add_run("技术报告 · 2026 年 10 月")
set_run_font(run, BODY_FONT, 10, color=RGBColor(96, 96, 96))

# ---------------------------------------------------------------------------
# Abstract
# ---------------------------------------------------------------------------
h2("摘要", number=None)
abstract = (
    "本文系统介绍一个面向四川麻将「血战到底」的 AI 工程：从完整规则引擎、状态编码、"
    "多任务策略网络，到行为克隆、DAgger 数据聚合与 PPO 自博弈的完整训练链路。"
    "规则引擎覆盖 108 张牌、定缺、换三张、碰杠、抢杠胡、一炮多响与血战续打，"
    "并以 461 项自动化测试保证正确性。策略网络以 404 维可见信息为输入，"
    "共享主干后接出牌（27 维）、响应（4 维）与价值（1 维）三个任务头。"
    "训练分为三步：先在 5000 局规则 AI 自对弈数据上做行为克隆；"
    "再用 600 局 DAgger 让规则 AI 为策略实际访问到的状态打标签；"
    "最后用带对手池的 PPO 做自博弈微调。在 300 局、2v2 轮换的评测中，"
    "纯行为克隆相对规则 AI 为 −1.13 分/座，DAgger 将对差距缩小到 −0.69 分/座；"
    "而当前规模的 PPO（480 局）反而退化到 −1.04 分/座。"
    "实验表明：在麻将这类不完全信息、多人博弈中，数据分布与标签质量比网络容量更关键，"
    "强化学习需要足够大的采样规模才能稳定超过模仿学习。"
)
p = doc.add_paragraph()
p.paragraph_format.line_spacing = 1.5
p.paragraph_format.first_line_indent = Cm(0.74)
run = p.add_run(abstract)
set_run_font(run, BODY_FONT, 10)

kw = doc.add_paragraph()
kw.paragraph_format.space_before = Pt(4)
run = kw.add_run("关键词：")
set_run_font(run, HEAD_FONT, 10.5, bold=True)
run = kw.add_run("四川麻将；血战到底；行为克隆；DAgger；PPO；不完全信息博弈")
set_run_font(run, BODY_FONT, 10.5)

# ---------------------------------------------------------------------------
# 1 Introduction
# ---------------------------------------------------------------------------
h1("引言", "1")
body(
    "麻将是一类典型的不完全信息、多人、随机博弈：玩家只能看到自己的手牌与公开的牌河、副露，"
    "需要在巨大的动作空间中做长期决策。四川麻将「血战到底」在通用麻将规则之上增加了"
    "定缺、换三张、刮风下雨、一炮多响、胡牌后继续对战等机制，规则复杂度与策略深度都显著提高。"
)
body(
    "现有麻将 AI 研究多集中于日本立直麻将，代表性工作 Suphx 通过监督学习预训练、"
    "强化学习自博弈与全局奖励预测达到高水平；而四川麻将的开源实现往往只覆盖规则或简单的启发式策略，"
    "缺少「规则引擎—状态编码—学习型策略—离线评估」的完整闭环。本文的目标是补齐这一闭环，"
    "并给出可复现的工程与实验。"
)
body("本文的主要贡献如下：")
contrib = [
    "实现完整的血战到底规则引擎，覆盖定缺、换三张、明杠/暗杠/补杠、抢杠胡、一炮多响、"
    "血战续打与可配置番型计分，并以 461 项测试保证正确性；",
    "设计 404 维状态编码与「出牌 + 响应 + 价值」三头策略网络，严格隔离对手暗牌；",
    "构建行为克隆 → DAgger → PPO 的三阶段训练流水线，并给出离线评测协议；",
    "通过对照实验量化各阶段收益：DAgger 明显缓解分布偏移，而小规模 PPO 会破坏已有策略。",
]
for i, text in enumerate(contrib, 1):
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.line_spacing = 1.5
    run = p.add_run(text)
    set_run_font(run, BODY_FONT, 10.5)

# ---------------------------------------------------------------------------
# 2 Related work
# ---------------------------------------------------------------------------
h1("相关工作", "2")
body(
    "Suphx 是微软提出的立直麻将 AI，采用「监督学习预训练 + 强化学习自博弈 + 全局奖励预测」的范式，"
    "在 Tenhou 平台达到高段位水平。DouZero 针对斗地主提出深度蒙特卡洛方法，"
    "说明在多人不完全信息卡牌游戏中，自博弈与价值估计同样有效。AlphaZero 的通用自博弈框架"
    "则展示了搜索与策略网络结合的上限。"
)
body(
    "在方法层面，本文使用行为克隆（Behavior Cloning, BC）作为冷启动，"
    "DAgger 通过让专家为策略访问到的状态打标签来缓解分布偏移，"
    "PPO 则用于在自博弈中直接优化期望得分。"
)
body(
    "开源生态方面，多数项目面向立直麻将或国标麻将，且许可证差异较大（AGPL、无许可证等），"
    "难以直接用于商业产品。因此本文选择自研规则引擎，仅参考公开工作的算法思想。"
)

# ---------------------------------------------------------------------------
# 3 Rules and engine
# ---------------------------------------------------------------------------
h1("规则与引擎", "3")
h2("规则要点", "3.1")
body(
    "血战到底使用 108 张牌（万、筒、条各 1–9，每种 4 张，无风牌字牌），四人参与。"
    "开局每人定缺一门，胡牌前必须打光该门；部分规则变体在定缺前增加「换三张」。"
    "对局中不能吃，只能碰、杠；杠后补摸一张并即时结算「刮风下雨」。"
    "一名玩家胡牌后不结束，胡牌者退出，其余人继续，直到只剩一人或牌墙摸完；"
    "同一张弃牌可被多家同时胡（一炮多响），补杠时其他玩家可以抢杠胡。"
)
caption("表 1  规则覆盖情况", above=True)
table(
    ["规则", "状态"],
    [
        ["108 张、无字牌、四人", "已实现"],
        ["定缺（缺一门）", "已实现"],
        ["换三张（可选开局）", "已实现"],
        ["碰 / 明杠 / 暗杠 / 补杠", "已实现"],
        ["抢杠胡", "已实现"],
        ["一炮多响", "已实现"],
        ["血战到底（胡牌者退出续打）", "已实现"],
        ["番型计分（可配置）", "已实现"],
        ["流局查花猪 / 查大叫", "简化版"],
        ["连庄 / 多局积分", "未实现"],
    ],
    widths=[9.0, 4.0],
)

h2("引擎架构", "3.2")
body(
    "引擎采用「反应式状态机」：调用方通过 legal_actions() 取得当前玩家的合法动作，"
    "用 step(action) 推进对局。状态机包含换三张、定缺、出牌、响应与结束五个阶段，"
    "所有规则、结算与血战流程由引擎统一保证，策略层无法绕过。图 1 给出系统分层："
    "规则层是唯一真值，决策层只消费编码后的可见信息，服务层负责对局编排与并发。"
)
figure("fig1_architecture.png", "图 1  系统总体架构", 15.0)

h2("正确性验证", "3.3")
body(
    "规则引擎以 pytest 风格的纯脚本测试覆盖牌型、向听、胡牌、番型与完整对局，"
    "共 461 项检查，其中包括 120 局随机对局与 40 局规则 AI 对局，并校验每局得分为零和。"
    "向听计算在纯 Python 下约 4400 手/秒，满足离线采样与实时决策的原型需求。"
)

# ---------------------------------------------------------------------------
# 4 State and action
# ---------------------------------------------------------------------------
h1("状态与动作表示", "4")
h2("状态特征", "4.1")
body(
    "状态编码只使用真实玩家可见的信息：自己的手牌与副露、四家的牌河与副露、四家的定缺、"
    "牌墙剩余、分数、已胡标记、刚摸到的牌，以及在响应阶段待处理的弃牌与出牌者。"
    "对手暗牌从不进入特征，编码器测试中专门验证了「替换对手手牌后特征不变」。"
    "表 2 列出 404 维特征的构成。"
)
caption("表 2  状态特征构成（404 维）", above=True)
table(
    ["特征段", "维度", "说明"],
    [
        ["手牌", "27", "每种牌数量 / 4"],
        ["自己副露（碰 / 杠）", "54", "两个 27 维 0/1 掩码"],
        ["四家牌河", "108", "4 × 27，除以 4"],
        ["四家副露", "108", "4 × 27，除以 4"],
        ["定缺", "12", "4 家 × 3 one-hot"],
        ["牌墙剩余", "1", "除以 55"],
        ["分数", "4", "除以 50"],
        ["已胡标记", "4", "0/1"],
        ["刚摸的牌", "27", "one-hot"],
        ["阶段标记", "1", "是否响应阶段"],
        ["待响应的牌", "27", "one-hot"],
        ["出牌者相对座位", "4", "one-hot"],
        ["每种牌剩余张数", "27", "可见牌之外剩余 / 4"],
        ["合计", "404", "—"],
    ],
    widths=[5.5, 2.2, 6.3],
)

h2("动作空间", "4.2")
body(
    "动作被拆成两个头：出牌头在 27 种牌上做分类，响应用 4 维分类（胡 / 碰 / 杠 / 过）。"
    "定缺、换三张、自摸、暗杠与补杠仍由规则策略处理。两个分类头在 softmax 之前"
    "对非法动作施加掩码，保证模型输出始终合法。"
)

# ---------------------------------------------------------------------------
# 5 Model
# ---------------------------------------------------------------------------
h1("模型结构", "5")
body(
    "策略网络采用共享主干的轻量多层感知机：输入 404 维，经过 404→512、512→512、512→256 "
    "三层全连接（ReLU 激活），再分叉到三个任务头。出牌头输出 27 维 logits，响应头输出 4 维 logits，"
    "价值头输出 1 维标量，用于估计当前局面的最终得分。图 2 给出结构示意。"
)
figure("fig2_model.png", "图 2  策略网络结构（PolicyValueNet）", 12.0)
body(
    "训练目标由三部分组成：出牌交叉熵、响应交叉熵与价值回归，"
    "L = CE_discard + CE_call + λ·MSE(value, target)，其中 λ = 0.5。"
    "网络参数量约 61 万，在 CPU 上即可完成训练与推理。"
)

# ---------------------------------------------------------------------------
# 6 Training
# ---------------------------------------------------------------------------
h1("训练方法", "6")
h2("行为克隆", "6.1")
body(
    "第一阶段用规则 AI 自对弈生成监督数据。采样按每 1000 局分片、在独立进程中执行，"
    "避免单进程长时间运行带来的内存膨胀。5000 局共得到 304,219 条决策样本，"
    "其中出牌 264,096 条、响应 40,123 条；只保留「有真实选择」的响应点（合法响应不少于 2 个），"
    "过滤掉大量无意义的「过」。模型训练 12 个 epoch，验证集出牌一致率 0.779、响应一致率 0.962。"
)
h2("DAgger 数据聚合", "6.2")
body(
    "纯行为克隆的弱点在于分布偏移：模型自己做出碰、杠、胡等决定后会进入训练集未覆盖的状态。"
    "DAgger 让当前策略实际对局，并由规则 AI 对策略访问到的状态给出专家动作，"
    "把这些标注样本并入训练集。本文用 600 局生成 20,885 条专家标注样本，"
    "与原始数据合并后重新训练。图 3 展示完整训练流水线。"
)
figure("fig3_pipeline.png", "图 3  训练流水线", 16.0)

h2("PPO 自博弈", "6.3")
body(
    "第三阶段以 DAgger 模型为初始策略，用 PPO 在自博弈中优化最终得分。"
    "每局由 2 个策略座与 2 个对手座组成，对手从「规则 AI + 历史 checkpoint」组成的对手池中随机抽取；"
    "奖励为本局最终得分除以 50 并裁剪到 ±5。每次迭代采样 32 局，共 15 次迭代，"
    "得到 51,816 条 transition。表 3 列出 PPO 超参数。"
)
caption("表 3  PPO 超参数", above=True)
table(
    ["参数", "值", "参数", "值"],
    [
        ["迭代次数", "15", "每轮局数", "32"],
        ["PPO epochs", "4", "minibatch", "512"],
        ["学习率", "3e−5", "clip", "0.2"],
        ["价值系数", "0.5", "熵系数", "0.01"],
        ["奖励裁剪", "±5", "梯度裁剪", "0.5"],
    ],
    widths=[3.6, 3.4, 3.6, 3.4],
)

# ---------------------------------------------------------------------------
# 7 Experiments
# ---------------------------------------------------------------------------
h1("实验", "7")
h2("实验设置", "7.1")
body(
    "评测采用 2 个模型座对 2 个规则 AI 座的混合对局，每局轮换座位以抵消庄家与位置偏差，"
    "共 300 局。指标为「每座平均得分差」（模型座平均分减规则 AI 座平均分），"
    "0 表示与规则 AI 持平，负数表示更弱。所有模型使用同一批随机种子，保证可比性。"
)
h2("主要结果", "7.2")
caption("表 4  不同模型对规则 AI 的强度（300 局，2v2 轮换）", above=True)
table(
    ["模型", "每座平均分", "与基线差"],
    [
        ["规则 AI（锚点）", "+0.565", "—"],
        ["SL（纯行为克隆）", "−0.565", "−1.13"],
        ["SL + PPO（240 局）", "−0.463", "−0.93"],
        ["SL + DAgger（600 局）", "−0.345", "−0.69"],
        ["DAgger + PPO（480 局 + 对手池）", "−0.518", "−1.04"],
    ],
    widths=[7.0, 4.0, 3.5],
)
figure("fig5_results.png", "图 4  不同模型的强度对比", 12.0)
body(
    "DAgger 是当前最佳方案：相比纯行为克隆把差距从 −1.13 缩小到 −0.69 分/座，"
    "说明「专家标注策略访问到的状态」确实缓解了分布偏移。"
    "而 PPO 在该规模下反而退化到 −1.04，训练日志中也出现价值损失突增的不稳定信号。"
)
h2("学习曲线", "7.3")
body(
    "图 5 对比了纯行为克隆与 DAgger 聚合在训练过程中的验证集出牌一致率。"
    "两者在前 4 个 epoch 快速上升，随后进入平台期；DAgger 在第 10 个 epoch 达到 0.780，"
    "略高于纯行为克隆的 0.779，但差异很小——真正的收益体现在实战得分而非离线一致率上。"
)
figure("fig4_curves.png", "图 5  验证集出牌一致率随训练轮次的变化", 12.0)
h2("分析", "7.4")
body(
    "第一，离线一致率与实战强度并不同步。DAgger 的一致率只提升约 0.1 个百分点，"
    "但实战差距缩小了 0.44 分/座，说明模型在「自己的状态分布」上做出了更正确的选择。"
    "第二，小规模 PPO 会破坏已有策略：480 局、15 次迭代的数据量不足以稳定估计优势函数，"
    "策略在噪声梯度下偏离了 DAgger 的良好行为。第三，规则 AI 仍然强于所有学习型策略，"
    "这与「专家数据来自规则 AI、模型只学到其近似」的预期一致。"
)

# ---------------------------------------------------------------------------
# 8 Discussion
# ---------------------------------------------------------------------------
h1("讨论与局限", "8")
body(
    "当前系统仍有多项局限。其一，模型只学习出牌与响应，定缺、自摸、暗杠与补杠仍由规则处理，"
    "策略上限受规则策略约束。其二，训练数据全部来自规则 AI 自对弈，模型无法超越数据分布的天花板。"
    "其三，PPO 的采样效率有限：单条前向推理与纯 Python 规则计算使大规模自博弈成本较高。"
    "其四，各地血战规则在番值、封顶、杠结算与流局赔偿上差异较大，上线前必须按目标平台逐项校准。"
)
body(
    "此外，麻将类产品涉及合规风险。本文的工程仅用于技术研究与产品开发，"
    "商业落地前应确保不涉及真实货币、兑换或变相赌博，并完成法律与数据合规评估。"
)

# ---------------------------------------------------------------------------
# 9 Conclusion
# ---------------------------------------------------------------------------
h1("结论与展望", "9")
body(
    "本文完成了四川麻将血战到底 AI 的完整工程闭环：规则引擎、状态编码、多任务策略网络、"
    "行为克隆、DAgger 与 PPO 训练，以及可复现的离线评测。实验表明，"
    "在数据有限时，DAgger 式数据聚合比小规模强化学习更有效；"
    "要把学习型策略推到规则 AI 之上，需要更大的自博弈规模、更稳定的 PPO 实现（GAE、KL 早停、"
    "学习率衰减）以及更强的对手池。"
)
body(
    "后续工作包括：把定缺、自摸与杠纳入动作空间；将向听计算改为按花色预计算以提升采样吞吐；"
    "引入多轮 DAgger 与 league 训练；并将模型导出为 ONNX、按服务化方案上线灰度。"
)

# ---------------------------------------------------------------------------
# References
# ---------------------------------------------------------------------------
h1("参考文献", None)
refs = [
    "LI J, KOYAMADA S, YE Q, et al. Suphx: Mastering Mahjong with Deep Reinforcement Learning[J]. arXiv:2003.13590, 2020.",
    "ZHA D, LAI J, CAO Y, et al. DouZero: Mastering DouDizhu with Self-Play Deep Reinforcement Learning[C]//Proceedings of ICML. 2021.",
    "SILVER D, HUBERT T, SCHRITTWIESER J, et al. A general reinforcement learning algorithm that masters chess, shogi, and Go through self-play[J]. Science, 2018, 362(6419): 1140-1144.",
    "SCHULMAN J, WOLSKI F, DHARIWAL P, et al. Proximal Policy Optimization Algorithms[J]. arXiv:1707.06347, 2017.",
    "ROSS S, GORDON G J, BAGNELL D. A Reduction of Imitation Learning and Structured Prediction to No-Regret Online Learning[C]//Proceedings of AISTATS. 2011: 627-635.",
]
for i, ref in enumerate(refs, 1):
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.4
    p.paragraph_format.left_indent = Cm(0.74)
    p.paragraph_format.first_line_indent = Cm(-0.74)
    run = p.add_run(f"[{i}] {ref}")
    set_run_font(run, BODY_FONT, 9)

doc.save(OUT)
print("saved", OUT)

import shutil

ascii_copy = ROOT / "paper" / "report.docx"
shutil.copyfile(OUT, ascii_copy)
print("saved", ascii_copy)
